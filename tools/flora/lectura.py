"""Lectura de las planillas de origen a listas de diccionarios."""
import openpyxl

COLS_NOTAS = [
    "Fotos propias o naturalist",
    "dato random (aveces relacionado con la imagen anterior, aveces no)",
    "Intento de asociacion mental",
    "col_T",
]
_MAX_VACIAS = 500


def _fila_vacia(fila) -> bool:
    return not any(v is not None and str(v).strip() for v in fila)


def _leer(ws, corte_vacias: int | None) -> list[dict]:
    filas = ws.iter_rows(values_only=True)
    encabezado = [h if h is not None else "col_T" for h in next(filas)]
    salida, vacias = [], 0
    for i, fila in enumerate(filas, start=2):
        if _fila_vacia(fila):
            vacias += 1
            if corte_vacias and vacias >= corte_vacias:
                break
            continue
        vacias = 0
        d = dict(zip(encabezado, fila))
        d["_fila"] = i
        salida.append(d)
    return salida


def leer_base(ruta) -> list[dict]:
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    return _leer(wb["Hoja1"], _MAX_VACIAS)


def leer_prefactibilidad(ruta) -> list[dict]:
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    return _leer(wb.active, None)
